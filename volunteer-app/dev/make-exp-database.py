"""
Writes the mock Exp volunteer database as an .xlsx to import into Google Sheets.

  cd volunteer-app
  python3 dev/make-exp-database.py <APP_SECRET> <output.xlsx>

<APP_SECRET>: the value you'll set as the Exp project's APP_SECRET script property
BEFORE running setup() (the demo accounts' passwords are stored under it).
Needs Node.js and Python's openpyxl.
"""
import json, subprocess, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

secret, out = sys.argv[1], sys.argv[2]
tabs = json.loads(subprocess.run(['node', 'dev/make-exp-database.mjs', secret], capture_output=True, text=True, check=True).stdout)

def clean(name, rows):
    head, body = rows[0], rows[1:]
    if name == 'Volunteer Guide':          # files lived in the demo's stand-in Google Drive
        t = head.index('Type'); body = [r for r in body if r[t] == 'text']
    if name == 'Trainee Photos':           # demo photos: stand-in Drive file IDs
        body = []
    if name == 'Volunteers':               # demo contacts: stand-in Google Contacts IDs
        for col in ('Google Contact ID', 'Google Contact'):
            if col in head:
                c = head.index(col)
                for r in body: r[c] = ''
    return [head] + body

wb = Workbook(); wb.remove(wb.active)
bold, fill = Font(bold=True), PatternFill('solid', fgColor='EEF2FF')
for t in tabs:
    rows = clean(t['name'], t['values'])
    ws = wb.create_sheet(t['name'])
    for r in rows:
        ws.append(['' if v is None else v for v in r])
    for cell in ws[1]:
        cell.font, cell.fill = bold, fill
    ws.freeze_panes = 'A2'
    for i, col in enumerate(rows[0], 1):
        width = max([len(str(col))] + [len(str(r[i - 1])) for r in rows[1:30] if i - 1 < len(r)])
        ws.column_dimensions[get_column_letter(i)].width = min(45, max(10, width + 2))
    # keep text that looks like numbers or dates as text (phones, NRICs, IDs, dates)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value, str):
                cell.number_format = '@'
    if t['hidden']:
        ws.sheet_state = 'hidden'
wb.save(out)
print(f'{out}: ' + ', '.join(f"{t['name']} ({len(clean(t['name'], t['values'])) - 1})" for t in tabs))
